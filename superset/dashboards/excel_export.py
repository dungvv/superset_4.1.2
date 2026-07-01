# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from superset.dashboards.word_export import get_image_stream, normalize_chart_name

PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*(.*?)\s*\}\}")
XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
DEFAULT_COLUMN_WIDTH = 8.43
DEFAULT_ROW_HEIGHT = 15
IMAGE_PADDING_PX = 8


def build_default_excel_workbook(
    charts: list[dict[str, Any]],
    dashboard_title: str,
    template_context: dict[str, Any] | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"

    ws["A1"] = dashboard_title
    context = template_context or {}
    if context.get("start_date") or context.get("end_date"):
        ws["A2"] = (
            f"{context.get('start_date', '')} - {context.get('end_date', '')}"
        )

    max_col = max((int(chart.get("col", 0)) for chart in charts), default=0) + 1
    start_row = 4
    block_height = 18
    block_width = 7

    for col_idx in range(1, max_col * block_width + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14

    for chart in charts:
        row = int(chart.get("row", 0))
        col = int(chart.get("col", 0))
        anchor_row = start_row + row * block_height
        anchor_col = 1 + col * block_width
        anchor = f"{get_column_letter(anchor_col)}{anchor_row}"
        target_range = CellRange(
            min_col=anchor_col,
            min_row=anchor_row,
            max_col=anchor_col + block_width - 1,
            max_row=anchor_row + block_height - 2,
        )

        ws[anchor] = chart.get("name") or ""
        add_chart_image_to_sheet(ws, chart, anchor, target_range)

    return save_workbook(wb)


def build_template_excel_workbook(
    charts: list[dict[str, Any]],
    template_path: str | Path,
    template_context: dict[str, Any] | None = None,
) -> bytes:
    wb = load_workbook(str(template_path))
    chart_by_name = {
        normalize_chart_name(chart.get("name")): chart
        for chart in charts
        if chart.get("name")
    }
    text_by_placeholder = {
        str(key).strip(): "" if value is None else str(value)
        for key, value in (template_context or {}).items()
    }

    for ws in wb.worksheets:
        replace_placeholders_in_worksheet(ws, chart_by_name, text_by_placeholder)

    return save_workbook(wb)


def replace_placeholders_in_worksheet(
    ws: Worksheet,
    chart_by_name: dict[str, dict[str, Any]],
    text_by_placeholder: dict[str, str],
) -> int:
    replacements = 0

    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str):
                continue

            cell_text = cell.value
            matches = list(PLACEHOLDER_PATTERN.finditer(cell_text))
            if not matches:
                continue

            exact_placeholder = (
                len(matches) == 1
                and matches[0].group(0).strip() == cell_text.strip()
            )
            if exact_placeholder:
                placeholder = matches[0].group(1).strip()
                chart = chart_by_name.get(normalize_chart_name(placeholder))
                if chart:
                    cell.value = None
                    add_chart_image_to_sheet(
                        ws,
                        chart,
                        cell.coordinate,
                        get_cell_range(ws, cell.coordinate),
                    )
                    replacements += 1
                    continue

            next_value = cell_text
            for match in matches:
                placeholder = match.group(1).strip()
                if placeholder in text_by_placeholder:
                    next_value = next_value.replace(
                        match.group(0),
                        text_by_placeholder[placeholder],
                    )

            if next_value != cell_text:
                cell.value = next_value
                replacements += 1

    return replacements


def get_cell_range(ws: Worksheet, coordinate: str) -> CellRange:
    for merged_range in ws.merged_cells.ranges:
        if coordinate in merged_range:
            return CellRange(str(merged_range))

    return CellRange(coordinate)


def add_chart_image_to_sheet(
    ws: Worksheet,
    chart: dict[str, Any],
    anchor: str,
    target_range: CellRange,
) -> bool:
    image_stream = get_image_stream(chart)
    if image_stream is None:
        return False

    image = ExcelImage(image_stream)
    target_width, target_height = get_range_size_px(ws, target_range)
    fit_image_to_box(
        image,
        max(target_width - IMAGE_PADDING_PX, 1),
        max(target_height - IMAGE_PADDING_PX, 1),
        int(chart.get("width", 0) or 0),
        int(chart.get("height", 0) or 0),
    )
    ws.add_image(image, anchor)
    return True


def get_range_size_px(ws: Worksheet, cell_range: CellRange) -> tuple[int, int]:
    min_col, min_row, max_col, max_row = range_boundaries(str(cell_range))
    width = 0
    height = 0

    for col_idx in range(min_col, max_col + 1):
        letter = get_column_letter(col_idx)
        column_width = ws.column_dimensions[letter].width or DEFAULT_COLUMN_WIDTH
        width += excel_column_width_to_px(float(column_width))

    for row_idx in range(min_row, max_row + 1):
        row_height = ws.row_dimensions[row_idx].height or DEFAULT_ROW_HEIGHT
        height += points_to_px(float(row_height))

    return width, height


def excel_column_width_to_px(width: float) -> int:
    return int(width * 7 + 5)


def points_to_px(points: float) -> int:
    return int(points * 96 / 72)


def fit_image_to_box(
    image: ExcelImage,
    max_width: int,
    max_height: int,
    source_width: int,
    source_height: int,
) -> None:
    width = source_width or image.width
    height = source_height or image.height

    if width <= 0 or height <= 0:
        return

    scale = min(max_width / width, max_height / height, 1)
    image.width = int(width * scale)
    image.height = int(height * scale)


def save_workbook(wb: Workbook) -> bytes:
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
