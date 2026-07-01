# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.
from __future__ import annotations

import base64
import io
import re
from copy import deepcopy
from pathlib import Path
from typing import Any, Iterator, NamedTuple

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches
from docx.table import Table
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run

PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*(.*?)\s*\}\}")
DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
WORD_EXPORT_PAGE_WIDTH = Cm(21.0)
WORD_EXPORT_PAGE_HEIGHT = Cm(100.0)
WORD_EXPORT_TOP_BOTTOM_MARGIN = Inches(0.2)
WORD_EXPORT_LEFT_RIGHT_MARGIN = Inches(0.25)
WORD_EXPORT_HEADER_FOOTER_DISTANCE = Inches(0)
WORD_EXPORT_MAX_IMAGE_WIDTH = Inches(6.8)


class ParagraphContext(NamedTuple):
    paragraph: Paragraph
    table: Table | None = None
    cell: Any | None = None


def normalize_chart_name(value: Any) -> str:
    return " ".join(str(value or "").split()).casefold()


def is_template_dashboard(
    dashboard_id: int | str,
    dashboard_slug: str | None,
    dashboard_title: str | None,
    configured_dashboards: Any,
) -> bool:
    if not configured_dashboards:
        return False

    if isinstance(configured_dashboards, str):
        configured_values = {
            item.strip()
            for item in configured_dashboards.split(",")
            if item.strip()
        }
    else:
        configured_values = {
            str(item).strip() for item in configured_dashboards if item
        }

    candidates = {
        str(dashboard_id),
        str(dashboard_slug or ""),
        str(dashboard_title or ""),
    }
    normalized_candidates = {normalize_chart_name(value) for value in candidates}
    normalized_configured_values = {
        normalize_chart_name(value) for value in configured_values
    }

    return bool(normalized_candidates & normalized_configured_values)


def apply_single_long_page_layout(doc: DocxDocument) -> None:
    for section in doc.sections:
        section.page_width = WORD_EXPORT_PAGE_WIDTH
        section.page_height = WORD_EXPORT_PAGE_HEIGHT
        section.top_margin = WORD_EXPORT_TOP_BOTTOM_MARGIN
        section.bottom_margin = WORD_EXPORT_TOP_BOTTOM_MARGIN
        section.left_margin = WORD_EXPORT_LEFT_RIGHT_MARGIN
        section.right_margin = WORD_EXPORT_LEFT_RIGHT_MARGIN
        section.header_distance = WORD_EXPORT_HEADER_FOOTER_DISTANCE
        section.footer_distance = WORD_EXPORT_HEADER_FOOTER_DISTANCE

    for context in iter_document_paragraphs(doc):
        context.paragraph.paragraph_format.page_break_before = False


def remove_manual_page_breaks(doc: DocxDocument) -> None:
    for context in iter_document_paragraphs(doc):
        paragraph = context.paragraph
        paragraph.paragraph_format.page_break_before = False
        page_breaks = paragraph._p.xpath(
            './/*[local-name()="br" and @*[local-name()="type"]="page"]',
        )
        rendered_page_breaks = paragraph._p.xpath(
            './/*[local-name()="lastRenderedPageBreak"]',
        )
        for element in list(page_breaks) + list(rendered_page_breaks):
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)


def build_default_word_document(
    charts: list[dict[str, Any]],
    dashboard_title: str,
) -> bytes:
    doc = Document()
    apply_single_long_page_layout(doc)

    title_para = doc.add_heading(dashboard_title, level=1)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    max_row = max(chart.get("row", 0) for chart in charts) + 1
    max_col = max(chart.get("col", 0) for chart in charts) + 1

    grid: list[list[list[dict[str, Any]] | None]] = [
        [None] * max_col for _ in range(max_row)
    ]
    for chart in charts:
        row = chart.get("row", 0)
        col = chart.get("col", 0)
        if 0 <= row < max_row and 0 <= col < max_col:
            if grid[row][col] is None:
                grid[row][col] = []
            grid[row][col].append(chart)

    table = doc.add_table(rows=0, cols=max_col)
    table.style = "Table Grid"
    table.autofit = True

    for grid_row in grid:
        cells = table.add_row().cells
        for col_idx, charts_in_cell in enumerate(grid_row):
            if not charts_in_cell:
                continue

            cell_para = cells[col_idx].paragraphs[0]
            for chart in charts_in_cell:
                image_stream = get_image_stream(chart)
                if image_stream is None:
                    continue

                run = cell_para.add_run()
                add_chart_picture(run, chart, image_stream)

    remove_manual_page_breaks(doc)
    return save_document(doc)


def build_template_word_document(
    charts: list[dict[str, Any]],
    template_path: str | Path,
    template_context: dict[str, Any] | None = None,
) -> bytes:
    doc = Document(str(template_path))
    apply_single_long_page_layout(doc)
    chart_by_name = {
        normalize_chart_name(chart.get("name")): chart
        for chart in charts
        if chart.get("name")
    }
    text_by_placeholder = {
        str(key).strip(): "" if value is None else str(value)
        for key, value in (template_context or {}).items()
    }

    replace_template_placeholders(doc, chart_by_name, text_by_placeholder)
    remove_manual_page_breaks(doc)
    return save_document(doc)


def replace_template_placeholders(
    doc: DocxDocument,
    chart_by_name: dict[str, dict[str, Any]],
    text_by_placeholder: dict[str, str] | None = None,
) -> int:
    replacements = 0
    for context in iter_document_paragraphs(doc):
        replacements += replace_template_placeholders_in_paragraph(
            context.paragraph,
            chart_by_name,
            text_by_placeholder or {},
            context.table,
            context.cell,
        )
    return replacements


def _run_has_image(run: Run) -> bool:
    """Return True if the run contains a drawing or picture element."""
    r = run._element
    return (
        r.find(qn("w:drawing")) is not None
        or r.find(qn("w:pict")) is not None
    )


def _copy_run_formatting(src_run: Run, dst_run: Run) -> None:
    """Copy the <w:rPr> formatting element from src_run to dst_run."""
    src_rPr = src_run._element.find(qn("w:rPr"))
    if src_rPr is None:
        return
    dst_r = dst_run._element
    existing = dst_r.find(qn("w:rPr"))
    if existing is not None:
        dst_r.remove(existing)
    dst_r.insert(0, deepcopy(src_rPr))


def _get_format_at_char(
    char: int,
    format_map: list[tuple[int, int, Run]],
) -> Run | None:
    """Return the sample run whose range covers the given character offset."""
    for start, end, run in format_map:
        if start <= char < end:
            return run
    return None


def _insert_text_preserving_format(
    paragraph: Paragraph,
    text: str,
    start_char: int,
    format_map: list[tuple[int, int, Run]],
    fallback_run: Run | None,
) -> None:
    """Insert text into paragraph, splitting into runs when formatting changes.

    This ensures that if the text spans multiple original runs with different
    formatting (e.g. italic vs normal), each sub-range gets its own run with
    the correct format copied from the original.
    """
    if not text:
        return

    i = 0
    while i < len(text):
        run = _get_format_at_char(start_char + i, format_map)
        # Find the furthest point where the same format applies
        j = i + 1
        while j < len(text) and _get_format_at_char(start_char + j, format_map) is run:
            j += 1

        new_run = paragraph.add_run(text[i:j])
        _copy_run_formatting(run or fallback_run, new_run)
        i = j


def _insert_image(
    paragraph: Paragraph,
    chart: dict[str, Any],
    sample_run: Run | None,
    fallback_run: Run | None,
    cell_width: int | None,
) -> bool:
    image_stream = get_image_stream(chart)
    if image_stream is not None:
        new_run = paragraph.add_run()
        _copy_run_formatting(sample_run or fallback_run, new_run)
        add_chart_picture(new_run, chart, image_stream, width=cell_width)
        return True
    return False

def _paragraph_has_drawing(paragraph: Paragraph) -> bool:
    return bool(
        paragraph._p.xpath('.//*[local-name()="drawing" or local-name()="pict"]')
    )

def _paragraph_is_empty(paragraph: Paragraph) -> bool:
    return not paragraph.text.strip() and not _paragraph_has_drawing(paragraph)

def _remove_empty_cell_paragraphs(cell: Any | None, keep: Paragraph) -> None:
    if cell is None:
        return

    for paragraph in list(cell.paragraphs):
        if paragraph._p is keep._p or not _paragraph_is_empty(paragraph):
            continue
        paragraph._p.getparent().remove(paragraph._p)


def replace_template_placeholders_in_paragraph(
    paragraph: Paragraph,
    chart_by_name: dict[str, dict[str, Any]],
    text_by_placeholder: dict[str, str],
    table: Table | None = None,
    cell: Any | None = None,
) -> int:
    original_runs = list(paragraph.runs)

    # Identify image runs and their indices in the original run order
    image_runs: list[tuple[int, Run]] = [
        (i, run) for i, run in enumerate(original_runs) if _run_has_image(run)
    ]

    # Build a format map: list of (start_char, end_char, sample_run) from all
    # text runs so we can look up the exact formatting for any character.
    format_map: list[tuple[int, int, Run]] = []
    char_offset = 0
    for run in original_runs:
        if _run_has_image(run):
            continue
        run_text = run.text or ""
        if run_text:
            format_map.append((char_offset, char_offset + len(run_text), run))
            char_offset += len(run_text)

    total_text = "".join(
        run.text or "" for run in original_runs if not _run_has_image(run)
    )

    matches = list(PLACEHOLDER_PATTERN.finditer(total_text))
    if not matches:
        return 0

    replacements = 0
    for match in matches:
        placeholder = match.group(1).strip()
        if (
            normalize_chart_name(placeholder) in chart_by_name
            or placeholder in text_by_placeholder
        ):
            replacements += 1

    if not replacements:
        return 0

    if table is not None:
        fix_table_widths(table)

    cell_width = get_cell_image_width(cell)

    # Global fallback run (first text run in the paragraph)
    fallback_run = next(
        (run for run in original_runs if not _run_has_image(run) and run.text),
        None,
    )

    # No images -> simple clear + rebuild with per-char formatting
    if not image_runs:
        paragraph.clear()
        cursor = 0
        inserted_chart_image = False
        for match in matches:
            if match.start() > cursor:
                _insert_text_preserving_format(
                    paragraph,
                    total_text[cursor : match.start()],
                    cursor,
                    format_map,
                    fallback_run,
                )
            placeholder = match.group(1).strip()
            chart = chart_by_name.get(normalize_chart_name(placeholder))
            if chart:
                sample = _get_format_at_char(match.start(), format_map)
                inserted_chart_image = (
                    _insert_image(paragraph, chart, sample, fallback_run, cell_width)
                    or inserted_chart_image
                )
            elif placeholder in text_by_placeholder:
                _insert_text_preserving_format(
                    paragraph,
                    text_by_placeholder[placeholder],
                    match.start(),
                    format_map,
                    fallback_run,
                )
            else:
                _insert_text_preserving_format(
                    paragraph,
                    match.group(0),
                    match.start(),
                    format_map,
                    fallback_run,
                )
            cursor = match.end()
        if cursor < len(total_text):
            _insert_text_preserving_format(
                paragraph,
                total_text[cursor:],
                cursor,
                format_map,
                fallback_run,
            )
        if inserted_chart_image:
            _remove_empty_cell_paragraphs(cell, paragraph)
        return replacements

    # Paragraph contains images. Build an ordered list of items that preserves
    # the original layout: image, text, image, text ... (or text, image, ...).
    items: list[tuple[str, Any]] = []
    prev_img_idx = -1

    for img_idx_actual, img_run in image_runs:
        seg_runs = [original_runs[i] for i in range(prev_img_idx + 1, img_idx_actual)]
        seg_text = "".join(r.text or "" for r in seg_runs)
        if seg_text:
            items.append(("text", seg_text))
        items.append(("image", img_run))
        prev_img_idx = img_idx_actual

    # Text after the last image
    seg_runs = [original_runs[i] for i in range(prev_img_idx + 1, len(original_runs))]
    seg_text = "".join(r.text or "" for r in seg_runs)
    if seg_text:
        items.append(("text", seg_text))

    # Remove all original runs
    for run in original_runs:
        paragraph._p.remove(run._element)

    # Rebuild paragraph item by item in the original order
    seg_char_offset = 0
    inserted_chart_image = False
    for item_type, item_value in items:
        if item_type == "image":
            paragraph._p.append(item_value._element)
            continue

        # Text item
        seg_text_local = item_value
        seg_start = seg_char_offset
        seg_end = seg_char_offset + len(seg_text_local)

        # Determine which matches fall inside this text segment
        seg_matches = [m for m in matches if seg_start <= m.start() < seg_end]

        cursor = seg_start
        for m in seg_matches:
            if m.start() > cursor:
                _insert_text_preserving_format(
                    paragraph,
                    total_text[cursor : m.start()],
                    cursor,
                    format_map,
                    fallback_run,
                )
            placeholder = m.group(1).strip()
            chart = chart_by_name.get(normalize_chart_name(placeholder))
            if chart:
                sample = _get_format_at_char(m.start(), format_map)
                inserted_chart_image = (
                    _insert_image(paragraph, chart, sample, fallback_run, cell_width)
                    or inserted_chart_image
                )
            elif placeholder in text_by_placeholder:
                _insert_text_preserving_format(
                    paragraph,
                    text_by_placeholder[placeholder],
                    m.start(),
                    format_map,
                    fallback_run,
                )
            else:
                _insert_text_preserving_format(
                    paragraph,
                    m.group(0),
                    m.start(),
                    format_map,
                    fallback_run,
                )
            cursor = m.end()
        if cursor < seg_end:
            _insert_text_preserving_format(
                paragraph,
                total_text[cursor : seg_end],
                cursor,
                format_map,
                fallback_run,
            )

        seg_char_offset += len(seg_text_local)

    if inserted_chart_image:
        _remove_empty_cell_paragraphs(cell, paragraph)

    return replacements


def iter_document_paragraphs(doc: DocxDocument) -> Iterator[ParagraphContext]:
    for paragraph in doc.paragraphs:
        yield ParagraphContext(paragraph)

    for table in doc.tables:
        yield from iter_table_paragraphs(table)

    for section in doc.sections:
        for story in (section.header, section.footer):
            for paragraph in story.paragraphs:
                yield ParagraphContext(paragraph)
            for table in story.tables:
                yield from iter_table_paragraphs(table)


def iter_table_paragraphs(table: Table) -> Iterator[ParagraphContext]:
    for row in table.rows:
        for cell in row.cells:
            yield from iter_cell_paragraphs(cell, table)


def iter_cell_paragraphs(
    cell: Any,
    table: Table,
) -> Iterator[ParagraphContext]:
    for paragraph in cell.paragraphs:
        yield ParagraphContext(paragraph, table, cell)

    for nested_table in cell.tables:
        yield from iter_table_paragraphs(nested_table)


def fix_table_widths(table: Table) -> None:
    table.autofit = False

    for column in table.columns:
        width = column.width
        if width:
            column.width = width

    for row in table.rows:
        for cell in row.cells:
            width = cell.width
            if width:
                cell.width = width


def get_cell_image_width(cell: Any | None) -> int | None:
    if cell is None:
        return None

    width = getattr(cell, "width", None)
    if not width:
        return None

    # Leave a small buffer for Word cell padding/borders so the image does not
    # force the column wider.
    return max(width - Inches(0.12), Inches(0.5))


def get_image_stream(chart: dict[str, Any]) -> io.BytesIO | None:
    image_data = chart.get("image", "")
    if not image_data.startswith("data:image/") or "," not in image_data:
        return None

    _header, encoded = image_data.split(",", 1)
    return io.BytesIO(base64.b64decode(encoded))


def add_chart_picture(
    run: Run,
    chart: dict[str, Any],
    image_stream: io.BytesIO,
    width: int | None = None,
) -> None:
    image_width = Inches(min(chart.get("width", 400) / 120, 4.5))

    if width is not None:
        image_width = min(width, WORD_EXPORT_MAX_IMAGE_WIDTH)
    else:
        image_width = min(image_width, WORD_EXPORT_MAX_IMAGE_WIDTH)

    run.add_picture(image_stream, width=image_width)


def save_document(doc: DocxDocument) -> bytes:
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()
